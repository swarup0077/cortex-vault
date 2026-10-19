"""Time a mixed demo upload end to end (SPEC §8 caps): copy into a session, phase 1, phase 2.

Generates a deterministic upload of 100 files per step (30 desktop screenshots, 10 phone screenshots, 25 camera
photos, 15 PDFs, 20 Word / PowerPoint / text files; `--files 200` doubles each kind with new seeds) and runs it
through `SessionManager.add_files`, the path a real upload takes after the browser has sent the files. Reports when
every file was listed (searchable by name and native text) and when every file was read (OCR, meaning, pictures),
with the stage times, files/sec per kind, and the session's size on disk (uploads, index copy, thumbnails). The
browser's transfer time isn't included: it depends on the visitor's connection, not the server. The caps are set
to fit the upload, so this measures the server whatever the deployed defaults are.

    python bench/bench_upload.py               # M5: 100 files, on the deploy host with nothing else busy
    python bench/bench_upload.py --files 200   # M6: the 200-file per-session cap
"""

from __future__ import annotations

import argparse
import json
import random
import shutil
import sys
import tempfile
import time
from dataclasses import replace
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(Path(__file__).resolve().parent)]

import docx  # noqa: E402
import pptx  # noqa: E402
import pymupdf  # noqa: E402

from bench_ocr import screenshot, sentence  # noqa: E402
from make_bench_corpus import page_image, phone, photo, png  # noqa: E402
from vault.config import Config  # noqa: E402
from vault.library import Library  # noqa: E402
from vault.session import SessionManager  # noqa: E402


def make_upload(out: Path, scale: int = 1) -> list[Path]:
    """100 × `scale` files in the M5 mix; scale 1 is exactly the M5 upload."""
    out.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []

    def add(path: Path) -> None:
        files.append(path)

    for i in range(30 * scale):
        p = out / f"Screenshot ({i + 1}).png"
        screenshot(700 + i)[0].save(p)
        add(p)
    for i in range(10 * scale):
        p = out / f"Screenshot_2026101{i % 10}-1015{i:02}_Chat.png"
        phone(800 + i, 2400).save(p)
        add(p)
    for i in range(25 * scale):
        p = out / f"IMG_{3000 + i}.jpg"
        photo(900 + i, sign=i < 5).save(p, quality=88)
        add(p)
    for i in range(10 * scale):  # PDFs with a text layer, 5 pages each
        rnd = random.Random(1000 + i)
        doc = pymupdf.open()
        for _ in range(5):
            page = doc.new_page()
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 790), " ".join(sentence(rnd, 12) for _ in range(30)),
                                fontsize=11)
        p = out / f"report {i + 1}.pdf"
        doc.save(p)
        add(p)
    for i in range(3 * scale):  # scanned, 3 pages, no text layer
        doc = pymupdf.open()
        for k in range(3):
            page = doc.new_page(width=612, height=792)
            page.insert_image(page.rect, stream=png(page_image(1100 + 10 * i + k)))
        p = out / f"scan {i + 1}.pdf"
        doc.save(p)
        add(p)
    for i in range(2 * scale):  # slide-style: picture bodies
        rnd = random.Random(1200 + i)
        doc = pymupdf.open()
        for k in range(4):
            page = doc.new_page(width=720, height=405)
            page.insert_text((40, 50), sentence(rnd, 6), fontsize=22)
            page.insert_image(pymupdf.Rect(40, 80, 680, 380), stream=png(page_image(1300 + 10 * i + k)),
                              keep_proportion=False)
        p = out / f"deck export {i + 1}.pdf"
        doc.save(p)
        add(p)
    for i in range(8 * scale):
        rnd = random.Random(1400 + i)
        d = docx.Document()
        for _ in range(12):
            d.add_paragraph(" ".join(sentence(rnd, 10) for _ in range(3)))
        p = out / f"notes {i + 1}.docx"
        d.save(str(p))
        add(p)
    for i in range(6 * scale):
        rnd = random.Random(1500 + i)
        prs = pptx.Presentation()
        for _ in range(8):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = sentence(rnd, 4)
            s.placeholders[1].text = sentence(rnd, 14)
        p = out / f"lecture {i + 1}.pptx"
        prs.save(str(p))
        add(p)
    for i in range(6 * scale):
        rnd = random.Random(1600 + i)
        p = out / (f"todo {i + 1}.txt" if i % 2 else f"readme {i + 1}.md")
        p.write_text("\n\n".join(" ".join(sentence(rnd, 9) for _ in range(4)) for _ in range(10)), encoding="utf-8")
        add(p)
    return files


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--files", type=int, default=100, help="a multiple of 100 (default 100, the M5 upload)")
    args = ap.parse_args()
    if args.files < 100 or args.files % 100:
        ap.error("--files must be a multiple of 100")
    work = Path(tempfile.mkdtemp(prefix="vault-upload-bench-"))
    try:
        source = make_upload(work / "source", args.files // 100)
        total_bytes = sum(p.stat().st_size for p in source)
        inbox = work / "gradio"
        for p in source:  # Gradio keeps each upload in its own folder; add_files moves them out
            (inbox / p.stem).mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, inbox / p.stem / p.name)
        files = [inbox / p.stem / p.name for p in source]
        largest = max(p.stat().st_size for p in source)
        base = Config(data_dir=ROOT / "demo_index")
        cfg = replace(base, upload_max_files=max(base.upload_max_files, len(source)),
                      upload_max_bytes=max(base.upload_max_bytes, largest),
                      upload_max_total=max(base.upload_max_total, total_bytes))
        demo = Library(cfg, read_only=True, base_dir=ROOT)
        sessions = SessionManager(demo, cfg, work / "sessions")
        marks: dict[str, float] = {}
        t0 = time.perf_counter()
        _, result = sessions.add_files(None, files, inbox, on_listed=lambda r: marks.setdefault(
            "listed", time.perf_counter() - t0))
        total = time.perf_counter() - t0
        rep = result.report
        on_disk = sum(f.stat().st_size for f in (work / "sessions").rglob("*") if f.is_file())
        print(json.dumps({
            "files": len(files), "megabytes": round(total_bytes / 2**20, 1),
            "accepted": sum(f.status == "indexed" for f in result.files),
            "not_accepted": {f.name: f.reason for f in result.files if f.status != "indexed"},
            "listed_s": round(marks.get("listed", 0.0), 1), "read_s": round(total, 1),
            "stages_s": {k: round(v, 1) for k, v in rep.timings.items()},
            "files_per_sec": {k: round(v, 2) for k, v in rep.files_per_sec().items()},
            "ocr_images": rep.ocr_images,
            "largest_mb": round(largest / 2**20, 1), "session_on_disk_mb": round(on_disk / 2**20, 1),
        }, indent=1))
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
