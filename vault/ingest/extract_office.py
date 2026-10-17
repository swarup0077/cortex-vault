"""Text from .docx, .pptx, .txt and .md (SPEC §6.3)."""

from __future__ import annotations

import re
import zipfile
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path

import docx
import pptx
from pptx.shapes.group import GroupShape

from vault.config import Config
from vault.ingest.records import Extracted, Segment, SkipFile


def zip_guard(path: Path, max_uncompressed: int, max_ratio: float) -> None:
    """Refuse zip bombs before python-docx/pptx unpack them."""
    with zipfile.ZipFile(path) as zf:
        infos = zf.infolist()
    total = sum(i.file_size for i in infos)
    packed = sum(i.compress_size for i in infos)
    if total > max_uncompressed:
        raise SkipFile(f"zip guard: {total / 2**20:.0f} MB uncompressed exceeds {max_uncompressed / 2**20:.0f} MB")
    if packed and total / packed > max_ratio:
        raise SkipFile(f"zip guard: compression ratio {total / packed:.0f} exceeds {max_ratio:.0f}")


def _iso(dt: datetime | None) -> str | None:
    return dt.replace(tzinfo=None).isoformat(timespec="seconds") if dt else None


def extract_docx(path: Path, cfg: Config) -> Extracted:
    """Paragraphs in order, then table cells row by row. Word stores no page numbers."""
    zip_guard(path, cfg.zip_max_uncompressed, cfg.zip_max_ratio)
    d = docx.Document(str(path))
    segments = [Segment(p.text, None, "doc_text") for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))  # merged cells repeat
            if cells:
                segments.append(Segment(" | ".join(cells), None, "doc_text"))
    return Extracted(segments=segments, doc_date=_iso(d.core_properties.created))


def _shape_texts(shapes) -> Iterator[str]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shape_texts(shape.shapes)
        elif getattr(shape, "has_text_frame", False):
            yield shape.text_frame.text
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield " | ".join(c.text for c in row.cells if c.text.strip())


def extract_pptx(path: Path, cfg: Config) -> Extracted:
    """Per slide: title, every text frame, tables, then speaker notes. page = slide number."""
    zip_guard(path, cfg.zip_max_uncompressed, cfg.zip_max_ratio)
    prs = pptx.Presentation(str(path))
    segments = []
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title  # proxies are rebuilt on access, so compare by shape_id
        parts = [title.text_frame.text] if title is not None else []
        title_id = title.shape_id if title is not None else None
        parts += _shape_texts(s for s in slide.shapes if s.shape_id != title_id)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text:
            segments.append(Segment(text, n, "slide_text"))
    return Extracted(segments=segments, pages=len(prs.slides), doc_date=_iso(prs.core_properties.created))


def extract_text(path: Path) -> Extracted:
    """.txt / .md as UTF-8 (errors replaced); paragraphs split on blank lines."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    paras = [p for p in re.split(r"\n\s*\n", raw) if p.strip()]
    return Extracted(segments=[Segment(p, None, "doc_text") for p in paras])
