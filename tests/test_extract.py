import os
import zipfile
from datetime import datetime
from pathlib import Path

import pptx
import pytest
from conftest import add_zip_member
from PIL import Image
from pptx.util import Inches

from vault.config import Config
from vault.ingest.extract_office import extract_docx, extract_pptx, extract_text, zip_guard
from vault.ingest.extract_pdf import extract_pdf, parse_pdf_date
from vault.ingest.records import SkipFile


def test_pdf_text_per_page(tmp_path: Path, make_pdf) -> None:
    p = make_pdf(tmp_path / "a b.pdf", ["Intro page", "", "Singular value decomposition is defined here"])
    ex = extract_pdf(p)
    assert ex.pages == 3
    assert [s.page for s in ex.segments] == [1, 3]  # blank page 2 yields nothing
    assert "Singular value decomposition" in ex.segments[1].text
    assert all(s.source == "pdf_text" for s in ex.segments)
    assert ex.doc_date == "2025-03-12T10:15:30"


def test_pdf_encrypted_is_skipped(tmp_path: Path, make_pdf) -> None:
    p = make_pdf(tmp_path / "locked.pdf", ["secret"], password="pw")
    with pytest.raises(SkipFile, match="encrypted"):
        extract_pdf(p)


def test_pdf_corrupt_raises(tmp_path: Path) -> None:
    p = tmp_path / "bad.pdf"
    p.write_bytes(b"not a pdf at all")
    with pytest.raises(Exception):
        extract_pdf(p)


@pytest.mark.parametrize(
    "raw,iso",
    [
        ("D:20250312101530+05'30'", "2025-03-12T10:15:30"),
        ("D:2024", "2024-01-01T00:00:00"),
        ("garbage", None),
        (None, None),
        ("D:20251399", None),
    ],
)
def test_parse_pdf_date(raw, iso) -> None:
    assert parse_pdf_date(raw) == iso


def test_docx_paragraphs_then_tables(tmp_path: Path, make_docx) -> None:
    p = make_docx(tmp_path / "n.docx", ["First para", "", "Second para"], table=[["k", "v"], ["a", "b"]])
    ex = extract_docx(p, Config(data_dir=tmp_path))
    assert [s.text for s in ex.segments] == ["First para", "Second para", "k | v", "a | b"]
    assert all(s.page is None and s.source == "doc_text" for s in ex.segments)


def test_pptx_per_slide(tmp_path: Path, make_pptx) -> None:
    p = make_pptx(
        tmp_path / "deck.pptx",
        [("Welcome", "agenda items", ""), ("Gradient descent", "step downhill", "say: learning rate")],
    )
    ex = extract_pptx(p, Config(data_dir=tmp_path))
    assert ex.pages == 2
    assert [s.page for s in ex.segments] == [1, 2]
    s1, s2 = (s.text for s in ex.segments)
    assert s1.startswith("Welcome") and "cell01" in s1 and s1.count("Welcome") == 1
    assert s2.splitlines() == ["Gradient descent", "step downhill", "say: learning rate"]


def test_text_paragraphs(tmp_path: Path) -> None:
    p = tmp_path / "n.md"
    p.write_bytes("# Title\n\nbody one\nstill one\n\n\n  \nbody two \xff".encode("latin-1"))
    ex = extract_text(p)
    assert [s.text for s in ex.segments][:2] == ["# Title", "body one\nstill one"]
    assert ex.segments[2].text.startswith("body two") and "�" in ex.segments[2].text


def test_zip_guard(tmp_path: Path, make_zip_bomb) -> None:
    bomb = make_zip_bomb(tmp_path / "bomb.docx")  # 50 MB of zeros as word/document.xml
    with pytest.raises(SkipFile, match="ratio"):
        zip_guard(bomb, 200 * 2**20, 100)
    with pytest.raises(SkipFile, match="XML parts"):
        zip_guard(bomb, 10 * 2**20, 1e9)
    with pytest.raises(SkipFile, match="uncompressed"):
        zip_guard(bomb, 200 * 2**20, 1e9, max_total=10 * 2**20)
    with pytest.raises(SkipFile):
        extract_docx(bomb, Config(data_dir=tmp_path))


def test_real_docx_passes_zip_guard(tmp_path: Path, make_docx) -> None:
    zip_guard(make_docx(tmp_path / "ok.docx", ["hello"]), 2**20, 100)


def _media_deck(path: Path) -> Path:
    """A deck with a real picture, a group shape holding a text box, notes, and 40 MB of incompressible video."""
    img = path.parent / "pic.png"
    Image.new("RGB", (64, 64), "red").save(img)
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Slippery slope"
    slide.shapes.add_picture(str(img), Inches(1), Inches(2))
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(4), Inches(2), Inches(3), Inches(1)).text_frame.text = "One step leads to another"
    slide.notes_slide.notes_text_frame.text = "Ask for examples"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank: no text
    prs.save(str(path))
    return add_zip_member(path, "ppt/media/media1.mp4", os.urandom(40 * 2**20))


def test_media_heavy_deck_is_indexed_without_reading_media(tmp_path: Path, monkeypatch) -> None:
    deck = _media_deck(tmp_path / "talk.pptx")
    read: list[str] = []
    real_read = zipfile.ZipFile.read
    monkeypatch.setattr(zipfile.ZipFile, "read", lambda self, name, pwd=None: read.append(name) or real_read(self, name, pwd))

    # The XML limit is far below the media size: media isn't counted against it.
    ex = extract_pptx(deck, Config(data_dir=tmp_path, zip_max_xml_bytes=2**20))
    assert ex.pages == 2
    assert [(s.page, s.text.splitlines()) for s in ex.segments] == [
        (1, ["Slippery slope", "One step leads to another", "Ask for examples"])
    ]
    assert read and not [n for n in read if "/media/" in n]


def test_demo_cap_still_refuses_big_decks(tmp_path: Path) -> None:
    deck = _media_deck(tmp_path / "talk.pptx")
    with pytest.raises(SkipFile, match="uncompressed"):
        extract_pptx(deck, Config(data_dir=tmp_path, zip_max_total=30 * 2**20))


def test_bomb_posing_as_media_is_refused(tmp_path: Path, make_pptx) -> None:
    deck = make_pptx(tmp_path / "d.pptx", [("t", "b", "")])
    add_zip_member(deck, "ppt/media/image9.png", bytes(20 * 2**20), compress=True)  # 20 MB of zeros
    with pytest.raises(SkipFile, match="image9.png compression ratio"):
        extract_pptx(deck, Config(data_dir=tmp_path))


def test_pptx_core_created_date(tmp_path: Path, make_pptx) -> None:
    p = make_pptx(tmp_path / "d.pptx", [("t", "b", "")])
    prs = pptx.Presentation(str(p))
    prs.core_properties.created = datetime(2025, 3, 12, 10, 15, 30)
    prs.save(str(p))
    assert extract_pptx(p, Config(data_dir=tmp_path)).doc_date == "2025-03-12T10:15:30"
