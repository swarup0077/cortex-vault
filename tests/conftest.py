"""Shared fixtures. All test files are generated here at test time; tests never touch personal files."""

from __future__ import annotations

import hashlib
import io
import re
import sqlite3
import zipfile
from collections.abc import Callable, Iterator, Sequence
from pathlib import Path

import docx
import numpy as np
import pptx
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx.util import Inches

from vault import db, models
from vault.config import Config
from vault.ingest.records import OcrLine

# ---------------------------------------------------------------- config / db


@pytest.fixture
def cfg(tmp_path: Path) -> Config:
    """Config with a throwaway data dir. The model cache stays shared so slow tests don't re-download."""
    return Config(data_dir=tmp_path / "data")


@pytest.fixture
def conn(cfg: Config) -> Iterator[sqlite3.Connection]:
    c = db.connect(cfg.db_path)
    yield c
    c.close()


# ---------------------------------------------------------------- models


class FakeTextEmbedder:
    """Deterministic hashed bag-of-words. Records every text it embeds so tests can count re-embeds."""

    name = "fake-hash-embedder"
    dim = 64

    def __init__(self) -> None:
        self.embedded: list[str] = []

    def _vec(self, text: str) -> np.ndarray:
        v = np.zeros(self.dim, np.float32)
        for tok in re.findall(r"\w+", text.lower()):
            v[int(hashlib.md5(tok.encode()).hexdigest(), 16) % self.dim] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_documents(self, texts):
        self.embedded += list(texts)
        return np.stack([self._vec(t) for t in texts]) if texts else np.zeros((0, self.dim), np.float32)

    def embed_query(self, text: str) -> np.ndarray:
        return self._vec(text)


@pytest.fixture
def fake_embedder() -> Iterator[FakeTextEmbedder]:
    e = FakeTextEmbedder()
    models.set_text_embedder(e)
    yield e
    models.set_text_embedder(None)


class FakeOcrEngine:
    """Scripted OCR: `lines_for(image)` gives each image's lines (none by default). Records image sizes."""

    name = "fake-ocr"

    def __init__(self) -> None:
        self.lines_for: Callable[[np.ndarray], Sequence[OcrLine]] | None = None
        self.seen: list[tuple[int, int]] = []  # (width, height) of every image read

    def read(self, image: np.ndarray) -> list[OcrLine]:
        self.seen.append((image.shape[1], image.shape[0]))
        return list(self.lines_for(image)) if self.lines_for else []


@pytest.fixture(autouse=True)
def fake_ocr() -> Iterator[FakeOcrEngine]:
    """Every test gets a fake OCR engine, so quick tests never load RapidOCR. Ask for `real_ocr` to opt out."""
    e = FakeOcrEngine()
    models.set_ocr_engine(e)
    yield e
    models.set_ocr_engine(None)


_real_ocr: models.OcrEngine | None = None


@pytest.fixture
def real_ocr(fake_ocr) -> models.OcrEngine:
    """The real RapidOCR engine (loaded once per test session). Tests using it are `slow`."""
    global _real_ocr
    if _real_ocr is None:
        _real_ocr = models.RapidOcrEngine(Config())
    models.set_ocr_engine(_real_ocr)
    return _real_ocr


# ---------------------------------------------------------------- file builders


def font(size: int) -> ImageFont.FreeTypeFont:
    """Pillow's bundled font, so drawn fixtures look the same on every OS."""
    return ImageFont.load_default(size=size)


def draw_text_image(size: tuple[int, int], texts: Sequence[tuple[int, int, str]], font_size: int = 32,
                    *, mode: str = "RGB") -> Image.Image:
    """Black text at (x, y) on white (or on transparent, for mode='RGBA')."""
    bg = (0, 0, 0, 0) if mode == "RGBA" else "white"
    img = Image.new(mode, size, bg)
    d = ImageDraw.Draw(img)
    for x, y, text in texts:
        d.text((x, y), text, fill="black", font=font(font_size))
    return img


def text_bbox(xy: tuple[int, int], text: str, font_size: int = 32) -> tuple[float, float, float, float]:
    """Where `draw_text_image` puts the ink of `text`."""
    return ImageDraw.Draw(Image.new("RGB", (1, 1))).textbbox(xy, text, font=font(font_size))


def save_image(img: Image.Image, path: Path, **kwargs) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, **kwargs)
    return path


def png_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def build_pdf_with_image(path: Path, title: str, image: Image.Image,
                         rect: tuple[float, float, float, float] = (50, 150, 550, 450), rotate: int = 0) -> Path:
    """One page: native `title` text plus `image` placed at `rect` (a slide exported with a picture body)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    page = doc.new_page()
    if title:
        page.insert_text((50, 80), title, fontsize=18)
    page.insert_image(pymupdf.Rect(*rect), stream=png_bytes(image), keep_proportion=False)
    page.set_rotation(rotate)
    doc.save(path)
    doc.close()
    return path


def build_scanned_pdf(path: Path, images: Sequence[Image.Image]) -> Path:
    """Pages that are only a full-page picture: no text layer."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    for img in images:
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes(img))
    doc.save(path)
    doc.close()
    return path


def build_pdf(
    path: Path, pages: Sequence[str], *, password: str | None = None, created: str = "D:20250312101530"
) -> Path:
    """A PDF with `pages[i]` as the native text of page i+1."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    doc.set_metadata({"creationDate": created})
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    kwargs = {}
    if password:
        kwargs = dict(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)
    doc.save(path, **kwargs)
    doc.close()
    return path


def build_docx(path: Path, paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    d.save(str(path))
    return path


def build_pptx(path: Path, slides: Sequence[tuple[str, str, str]]) -> Path:
    """`slides` are (title, body, speaker notes). Slide 1 is also given a 2×2 table."""
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = pptx.Presentation()
    for i, (title, body, notes) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if i == 0:
            tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
            for r in range(2):
                for c in range(2):
                    tbl.cell(r, c).text = f"cell{r}{c}"
    prs.save(str(path))
    return path


def build_zip_bomb(path: Path, size: int = 50 * 2**20) -> Path:
    """A highly compressible archive with an Office extension."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"\0" * size)
    return path


def add_zip_member(path: Path, name: str, data: bytes, *, compress: bool = False) -> Path:
    """Append a part to an existing .docx/.pptx (e.g. a big media file, or a bomb posing as one)."""
    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED) as zf:
        zf.writestr(name, data)
    return path


@pytest.fixture
def make_pdf() -> Callable[..., Path]:
    return build_pdf


@pytest.fixture
def make_docx() -> Callable[..., Path]:
    return build_docx


@pytest.fixture
def make_pptx() -> Callable[..., Path]:
    return build_pptx


@pytest.fixture
def make_zip_bomb() -> Callable[..., Path]:
    return build_zip_bomb


# ---------------------------------------------------------------- a small mixed corpus for search tests


def build_search_corpus(root: Path) -> Path:
    """Targets for the M1 search criteria plus distractors. No query word appears in the docx target."""
    build_pdf(
        root / "papers" / "linear algebra notes.pdf",
        [
            "Course overview: vectors, matrices and linear maps between vector spaces.",
            "Definition. The singular value decomposition of a matrix A factors it as U Sigma V transpose, "
            "where U and V are orthogonal and Sigma is diagonal with non-negative entries.",
            "Exercises: compute the eigenvalues of the following matrices by hand.",
        ],
    )
    build_pdf(
        root / "papers" / "diffusion.pdf",
        ["Denoising diffusion probabilistic models generate images by reversing a noising process."],
    )
    build_pptx(
        root / "decks" / "ml lecture.pptx",
        [
            ("Welcome", "Agenda for today", ""),
            ("Loss functions", "Mean squared error and cross entropy", ""),
            ("Gradient descent", "Update the weights by stepping against the gradient", "learning rate sets the step"),
            ("Questions", "Thank you", ""),
        ],
    )
    build_docx(
        root / "home" / "household.docx",
        ["The electricity invoice from the utility company must be settled by the fifth of every month."],
    )
    build_docx(root / "home" / "travel.docx", ["Train to Goa leaves at 6 am; pack sunscreen and a hat."])
    (root / "home" / "recipe.txt").write_text(
        "Sourdough: mix flour, water and starter, then proof overnight.", encoding="utf-8"
    )
    return root


@pytest.fixture(scope="session")
def make_search_corpus() -> Callable[[Path], Path]:
    return build_search_corpus
