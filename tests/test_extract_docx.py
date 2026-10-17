"""The direct-XML .docx reader must match python-docx's text, without loading media."""

import os
import zipfile
from datetime import datetime
from pathlib import Path

import docx
from conftest import add_zip_member
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from PIL import Image

from vault.config import Config
from vault.ingest.extract_office import extract_docx


def python_docx_reference(path: Path) -> list[str]:
    """What M1's python-docx extractor produced: body paragraphs, then table rows."""
    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
            if cells:
                out.append(" | ".join(cells))
    return out


def rich_docx(path: Path) -> Path:
    """Tabs, line and page breaks, horizontal and vertical merges, a picture and 40 MB of media."""
    img = path.parent / "pic.png"
    Image.new("RGB", (32, 32), "blue").save(img)
    d = docx.Document()
    d.add_paragraph("Quarterly report")
    p = d.add_paragraph("Name:")
    run = p.add_run()
    run.add_tab()
    run.add_text("Ada")
    run.add_break()  # line break: newline
    run.add_text("second line")
    run.add_break(WD_BREAK.PAGE)  # page break: no text
    run.add_text("after page break")
    d.add_paragraph("")
    d.add_picture(str(img), width=Inches(1))
    t = d.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"r{r}c{c}"
    t.cell(0, 0).merge(t.cell(0, 1))  # horizontal
    t.cell(1, 2).merge(t.cell(2, 2))  # vertical
    t.cell(2, 0).text = "same"
    t.cell(2, 1).text = "same"  # equal texts in one row collapse, as before
    d.add_paragraph("Closing paragraph after the table")
    d.core_properties.created = datetime(2025, 3, 12, 10, 15, 30)
    d.save(str(path))
    return add_zip_member(path, "word/media/video1.mp4", os.urandom(40 * 2**20))


def test_matches_python_docx_on_rich_fixture(tmp_path: Path) -> None:
    p = rich_docx(tmp_path / "report.docx")
    ex = extract_docx(p, Config(data_dir=tmp_path))
    assert [s.text for s in ex.segments] == python_docx_reference(p)
    assert ex.segments[1].text == "Name:\tAda\nsecond lineafter page break"
    assert ex.doc_date == "2025-03-12T10:15:30"


def test_matches_python_docx_on_simple_fixture(tmp_path: Path, make_docx) -> None:
    p = make_docx(tmp_path / "n.docx", ["First para", "", "Second para"], table=[["k", "v"], ["a", "b"]])
    assert [s.text for s in extract_docx(p, Config(data_dir=tmp_path)).segments] == python_docx_reference(p)


def test_media_is_never_read(tmp_path: Path, monkeypatch) -> None:
    p = rich_docx(tmp_path / "report.docx")
    read: list[str] = []
    real_read = zipfile.ZipFile.read
    monkeypatch.setattr(zipfile.ZipFile, "read", lambda self, name, pwd=None: read.append(name) or real_read(self, name, pwd))
    extract_docx(p, Config(data_dir=tmp_path, zip_max_xml_bytes=2**20))
    assert read and not [n for n in read if "/media/" in n]
